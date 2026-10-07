from fastapi import FastAPI, APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import io
import re
import logging
import uuid
from pathlib import Path
from datetime import datetime, timezone
from typing import List, Optional, Dict

from pydantic import BaseModel, Field

from money import parse_amount_to_cents, format_cents
import ocr as ocr_engine

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

app = FastAPI()
api_router = APIRouter(prefix="/api")

logging.basicConfig(level=logging.INFO,
                    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

MONTHS_IT = ["Gennaio", "Febbraio", "Marzo", "Aprile", "Maggio", "Giugno",
             "Luglio", "Agosto", "Settembre", "Ottobre", "Novembre", "Dicembre"]

DEFAULT_Z_ALIASES = [
    "TOT. CORRISPETTIVO GIORNO", "TOT CORRISPETTIVO GIORNO",
    "TOTALE CORRISPETTIVO GIORNO", "TOTALE CORRISPETTIVI", "CORRISPETTIVO GIORNO",
]

STATUS_VERIFICATO = "VERIFICATO"
STATUS_NON_QUADRATO = "NON_QUADRATO"
STATUS_DA_CONTROLLARE = "DA_CONTROLLARE"


# --------------------------------------------------------------------------- #
# Models                                                                       #
# --------------------------------------------------------------------------- #
def now_iso():
    return datetime.now(timezone.utc).isoformat()


class ColumnIn(BaseModel):
    display_name: str
    ocr_label: str = ""
    aliases: List[str] = []
    include_in_reconciliation: bool = True
    active: bool = True


class ColumnUpdate(BaseModel):
    display_name: Optional[str] = None
    ocr_label: Optional[str] = None
    aliases: Optional[List[str]] = None
    include_in_reconciliation: Optional[bool] = None
    active: Optional[bool] = None


class RegisterIn(BaseModel):
    name: str
    business_name: str = ""
    year: Optional[int] = None
    start_date: Optional[str] = None
    end_date: Optional[str] = None
    total_aliases: List[str] = Field(default_factory=lambda: list(DEFAULT_Z_ALIASES))
    columns: List[ColumnIn] = []


class RegisterUpdate(BaseModel):
    name: Optional[str] = None
    business_name: Optional[str] = None
    year: Optional[int] = None
    start_date: Optional[str] = None
    end_date: Optional[str] = None
    total_aliases: Optional[List[str]] = None


class ReorderIn(BaseModel):
    column_ids: List[str]


class ZReportIn(BaseModel):
    id: Optional[str] = None
    amount_cents: Optional[int] = None
    detected_date: Optional[str] = None
    confidence: Optional[float] = None


class EntryIn(BaseModel):
    business_date: str  # YYYY-MM-DD
    total_receipts_cents: Optional[int] = None
    z_reports: Optional[List[ZReportIn]] = None  # azzeramenti della giornata
    values: Dict[str, Optional[int]] = {}  # column_id -> cents
    printed_total_cents: Optional[int] = None
    status_hint: Optional[str] = None  # DA_CONTROLLARE to force
    ocr_raw: Optional[dict] = None
    replace_id: Optional[str] = None


class OCRRequest(BaseModel):
    image_base64: str


# --------------------------------------------------------------------------- #
# Helpers                                                                      #
# --------------------------------------------------------------------------- #
def clean(doc: dict) -> dict:
    doc.pop("_id", None)
    return doc


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip().upper())


async def get_register_or_404(register_id: str) -> dict:
    reg = await db.registers.find_one({"id": register_id}, {"_id": 0})
    if not reg:
        raise HTTPException(status_code=404, detail="Registro non trovato")
    return reg


async def get_columns(register_id: str) -> List[dict]:
    cols = await db.columns.find({"register_id": register_id}, {"_id": 0}).to_list(1000)
    cols.sort(key=lambda c: c.get("order", 0))
    return cols


def compute_reconciliation(columns: List[dict], total_receipts_cents, values: Dict[str, Optional[int]], status_hint):
    included_ids = {c["id"] for c in columns if c.get("include_in_reconciliation") and c.get("active", True)}
    departments_total = 0
    for cid, amt in (values or {}).items():
        if cid in included_ids and amt is not None:
            departments_total += int(amt)
    total = total_receipts_cents if total_receipts_cents is not None else None
    difference = None
    if total is not None:
        difference = total - departments_total
    if status_hint == STATUS_DA_CONTROLLARE:
        status = STATUS_DA_CONTROLLARE
    elif total is None:
        status = STATUS_DA_CONTROLLARE
    elif difference == 0:
        status = STATUS_VERIFICATO
    else:
        status = STATUS_NON_QUADRATO
    return departments_total, difference, status


def build_zreports(items: List[ZReportIn], existing: Optional[List[dict]] = None) -> List[dict]:
    """Normalize incoming azzeramenti into child subdocuments with stable ids,
    order_index and timestamps. Preserves created_at for existing ids."""
    existing_map = {z.get("id"): z for z in (existing or []) if z.get("id")}
    out = []
    for i, it in enumerate(items):
        zid = it.id or str(uuid.uuid4())
        prev = existing_map.get(zid) or {}
        out.append({
            "id": zid,
            "amount_cents": it.amount_cents,
            "detected_date": it.detected_date,
            "confidence": it.confidence,
            "order_index": i,
            "created_at": prev.get("created_at", now_iso()),
            "updated_at": now_iso(),
        })
    return out


def zreports_total(zreports: Optional[List[dict]]) -> Optional[int]:
    """Total corrispettivi = exact integer-cent sum of all azzeramenti of the day."""
    if not zreports:
        return None
    vals = [int(z["amount_cents"]) for z in zreports if z.get("amount_cents") is not None]
    if not vals:
        return None
    return sum(vals)


def map_lines_to_columns(lines: List[dict], columns: List[dict]):
    """Map OCR reparti lines to configured columns using ocr_label + aliases.
    Prefers the longest matching alias to avoid 'ALIMENTARI' catching 'ALIMENTARI 22'."""
    values: Dict[str, Optional[int]] = {}
    mapping = []  # per line: matched column id/name
    unmatched = []
    for ln in lines:
        label_n = _norm(ln.get("label"))
        best_col = None
        best_score = 0
        for col in columns:
            if not col.get("active", True):
                continue
            candidates = [col.get("ocr_label", "")] + (col.get("aliases") or [])
            for cand in candidates:
                cand_n = _norm(cand)
                if not cand_n:
                    continue
                score = 0
                if label_n == cand_n:
                    score = len(cand_n) + 100
                elif label_n.startswith(cand_n + " ") or (" " + cand_n) in (" " + label_n):
                    score = len(cand_n)
                elif cand_n in label_n:
                    score = len(cand_n) - 1
                if score > best_score:
                    best_score = score
                    best_col = col
        if best_col is not None:
            values[best_col["id"]] = ln.get("amount_cents")
            mapping.append({"label": ln.get("label"), "column_id": best_col["id"],
                            "column_name": best_col["display_name"],
                            "amount_cents": ln.get("amount_cents"),
                            "confidence": ln.get("confidence")})
        else:
            unmatched.append({"label": ln.get("label"), "amount_cents": ln.get("amount_cents"),
                              "confidence": ln.get("confidence")})
    return values, mapping, unmatched


# --------------------------------------------------------------------------- #
# Register routes                                                              #
# --------------------------------------------------------------------------- #
@api_router.get("/")
async def root():
    return {"message": "Registro Corrispettivi API"}


@api_router.get("/registers")
async def list_registers():
    regs = await db.registers.find({}, {"_id": 0}).to_list(1000)
    regs.sort(key=lambda r: r.get("updated_at", ""), reverse=True)
    for r in regs:
        r["entries_count"] = await db.entries.count_documents({"register_id": r["id"]})
    return regs


@api_router.post("/registers")
async def create_register(payload: RegisterIn):
    reg_id = str(uuid.uuid4())
    doc = {
        "id": reg_id,
        "name": payload.name,
        "business_name": payload.business_name,
        "year": payload.year,
        "start_date": payload.start_date,
        "end_date": payload.end_date,
        "total_aliases": payload.total_aliases or list(DEFAULT_Z_ALIASES),
        "created_at": now_iso(),
        "updated_at": now_iso(),
    }
    await db.registers.insert_one(dict(doc))
    for i, col in enumerate(payload.columns):
        await db.columns.insert_one({
            "id": str(uuid.uuid4()),
            "register_id": reg_id,
            "display_name": col.display_name,
            "ocr_label": col.ocr_label,
            "aliases": col.aliases,
            "order": i,
            "include_in_reconciliation": col.include_in_reconciliation,
            "active": col.active,
            "created_at": now_iso(),
        })
    return clean(doc)


@api_router.get("/registers/{register_id}")
async def get_register(register_id: str):
    reg = await get_register_or_404(register_id)
    reg["columns"] = await get_columns(register_id)
    return reg


@api_router.put("/registers/{register_id}")
async def update_register(register_id: str, payload: RegisterUpdate):
    await get_register_or_404(register_id)
    update = {k: v for k, v in payload.model_dump().items() if v is not None}
    update["updated_at"] = now_iso()
    await db.registers.update_one({"id": register_id}, {"$set": update})
    return await get_register(register_id)


@api_router.delete("/registers/{register_id}")
async def delete_register(register_id: str):
    await get_register_or_404(register_id)
    await db.registers.delete_one({"id": register_id})
    await db.columns.delete_many({"register_id": register_id})
    await db.entries.delete_many({"register_id": register_id})
    return {"ok": True}


@api_router.post("/registers/{register_id}/duplicate")
async def duplicate_register(register_id: str, payload: RegisterUpdate):
    src = await get_register_or_404(register_id)
    new_id = str(uuid.uuid4())
    doc = {
        "id": new_id,
        "name": payload.name or f"{src['name']} (copia)",
        "business_name": payload.business_name if payload.business_name is not None else src.get("business_name", ""),
        "year": payload.year if payload.year is not None else src.get("year"),
        "start_date": payload.start_date,
        "end_date": payload.end_date,
        "total_aliases": src.get("total_aliases", list(DEFAULT_Z_ALIASES)),
        "created_at": now_iso(),
        "updated_at": now_iso(),
    }
    await db.registers.insert_one(dict(doc))
    for col in await get_columns(register_id):
        await db.columns.insert_one({
            "id": str(uuid.uuid4()),
            "register_id": new_id,
            "display_name": col["display_name"],
            "ocr_label": col.get("ocr_label", ""),
            "aliases": col.get("aliases", []),
            "order": col.get("order", 0),
            "include_in_reconciliation": col.get("include_in_reconciliation", True),
            "active": col.get("active", True),
            "created_at": now_iso(),
        })
    # NOTE: entries and amounts are intentionally NOT copied
    return clean(doc)


# --------------------------------------------------------------------------- #
# Column routes                                                                #
# --------------------------------------------------------------------------- #
@api_router.get("/registers/{register_id}/columns")
async def list_columns(register_id: str):
    await get_register_or_404(register_id)
    return await get_columns(register_id)


@api_router.post("/registers/{register_id}/columns")
async def add_column(register_id: str, payload: ColumnIn):
    await get_register_or_404(register_id)
    existing = await get_columns(register_id)
    order = max([c.get("order", 0) for c in existing], default=-1) + 1
    doc = {
        "id": str(uuid.uuid4()),
        "register_id": register_id,
        "display_name": payload.display_name,
        "ocr_label": payload.ocr_label,
        "aliases": payload.aliases,
        "order": order,
        "include_in_reconciliation": payload.include_in_reconciliation,
        "active": payload.active,
        "created_at": now_iso(),
    }
    await db.columns.insert_one(dict(doc))
    await db.registers.update_one({"id": register_id}, {"$set": {"updated_at": now_iso()}})
    return clean(doc)


@api_router.put("/columns/{column_id}")
async def update_column(column_id: str, payload: ColumnUpdate):
    col = await db.columns.find_one({"id": column_id}, {"_id": 0})
    if not col:
        raise HTTPException(status_code=404, detail="Colonna non trovata")
    update = {k: v for k, v in payload.model_dump().items() if v is not None}
    await db.columns.update_one({"id": column_id}, {"$set": update})
    return clean(await db.columns.find_one({"id": column_id}))


@api_router.delete("/columns/{column_id}")
async def delete_column(column_id: str):
    col = await db.columns.find_one({"id": column_id}, {"_id": 0})
    if not col:
        raise HTTPException(status_code=404, detail="Colonna non trovata")
    await db.columns.delete_one({"id": column_id})
    # remove stored values referencing this column
    await db.entries.update_many({"register_id": col["register_id"]},
                                 {"$unset": {f"values.{column_id}": ""}})
    return {"ok": True}


@api_router.put("/registers/{register_id}/columns/reorder")
async def reorder_columns(register_id: str, payload: ReorderIn):
    await get_register_or_404(register_id)
    for i, cid in enumerate(payload.column_ids):
        await db.columns.update_one({"id": cid, "register_id": register_id}, {"$set": {"order": i}})
    await db.registers.update_one({"id": register_id}, {"$set": {"updated_at": now_iso()}})
    return await get_columns(register_id)


# --------------------------------------------------------------------------- #
# Entry routes                                                                 #
# --------------------------------------------------------------------------- #
@api_router.get("/registers/{register_id}/entries")
async def list_entries(register_id: str, month: Optional[str] = None):
    await get_register_or_404(register_id)
    query = {"register_id": register_id}
    if month:
        query["business_date"] = {"$regex": f"^{month}"}
    entries = await db.entries.find(query, {"_id": 0}).to_list(2000)
    entries.sort(key=lambda e: e.get("business_date", ""))
    return entries


@api_router.get("/registers/{register_id}/entries/check")
async def check_duplicate(register_id: str, date: str):
    existing = await db.entries.find_one({"register_id": register_id, "business_date": date}, {"_id": 0})
    return {"exists": existing is not None, "entry": existing}


@api_router.post("/registers/{register_id}/entries")
async def create_entry(register_id: str, payload: EntryIn):
    await get_register_or_404(register_id)
    columns = await get_columns(register_id)

    if payload.z_reports is not None:
        zreports = build_zreports(payload.z_reports)
        total = zreports_total(zreports)
    else:
        zreports = None
        total = payload.total_receipts_cents
    departments_total, difference, status = compute_reconciliation(
        columns, total, payload.values, payload.status_hint)

    existing = await db.entries.find_one(
        {"register_id": register_id, "business_date": payload.business_date}, {"_id": 0})

    if existing and not payload.replace_id:
        raise HTTPException(status_code=409,
                            detail=f"Esiste gia' una registrazione per il {payload.business_date}")

    base = {
        "register_id": register_id,
        "business_date": payload.business_date,
        "total_receipts_cents": total,
        "values": payload.values,
        "departments_total_cents": departments_total,
        "printed_total_cents": payload.printed_total_cents,
        "difference_cents": difference,
        "reconciliation_status": status,
        "ocr_raw": payload.ocr_raw,
        "updated_at": now_iso(),
    }
    if zreports is not None:
        base["z_reports"] = zreports
    if payload.replace_id and existing:
        await db.entries.update_one({"id": existing["id"]}, {"$set": base})
        entry_id = existing["id"]
    else:
        entry_id = str(uuid.uuid4())
        base["id"] = entry_id
        base["created_at"] = now_iso()
        if "z_reports" not in base:
            base["z_reports"] = []
        await db.entries.insert_one(dict(base))
    await db.registers.update_one({"id": register_id}, {"$set": {"updated_at": now_iso()}})
    return clean(await db.entries.find_one({"id": entry_id}))


@api_router.put("/entries/{entry_id}")
async def update_entry(entry_id: str, payload: EntryIn):
    entry = await db.entries.find_one({"id": entry_id}, {"_id": 0})
    if not entry:
        raise HTTPException(status_code=404, detail="Giornata non trovata")
    columns = await get_columns(entry["register_id"])

    if payload.z_reports is not None:
        zreports = build_zreports(payload.z_reports, existing=entry.get("z_reports"))
        total = zreports_total(zreports)
    else:
        zreports = entry.get("z_reports")
        total = zreports_total(zreports) if zreports else payload.total_receipts_cents
    departments_total, difference, status = compute_reconciliation(
        columns, total, payload.values, payload.status_hint)
    update = {
        "business_date": payload.business_date,
        "total_receipts_cents": total,
        "values": payload.values,
        "departments_total_cents": departments_total,
        "printed_total_cents": payload.printed_total_cents,
        "difference_cents": difference,
        "reconciliation_status": status,
        "updated_at": now_iso(),
    }
    if payload.z_reports is not None:
        update["z_reports"] = zreports
    if payload.ocr_raw is not None:
        update["ocr_raw"] = payload.ocr_raw
    await db.entries.update_one({"id": entry_id}, {"$set": update})
    return clean(await db.entries.find_one({"id": entry_id}))


@api_router.delete("/entries/{entry_id}")
async def delete_entry(entry_id: str):
    res = await db.entries.delete_one({"id": entry_id})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Giornata non trovata")
    return {"ok": True}


# --------------------------------------------------------------------------- #
# OCR routes                                                                   #
# --------------------------------------------------------------------------- #
@api_router.post("/registers/{register_id}/ocr/azzeramento")
async def ocr_azzeramento(register_id: str, payload: OCRRequest):
    reg = await get_register_or_404(register_id)
    try:
        result = await ocr_engine.parse_azzeramento(
            payload.image_base64, reg.get("total_aliases", DEFAULT_Z_ALIASES))
    except Exception as e:
        logger.exception("OCR azzeramento failed")
        raise HTTPException(status_code=502, detail=f"Errore OCR azzeramento: {e}")
    return result


@api_router.post("/registers/{register_id}/ocr/reparti")
async def ocr_reparti(register_id: str, payload: OCRRequest):
    await get_register_or_404(register_id)
    columns = await get_columns(register_id)
    try:
        result = await ocr_engine.parse_reparti(payload.image_base64)
    except Exception as e:
        logger.exception("OCR reparti failed")
        raise HTTPException(status_code=502, detail=f"Errore OCR reparti: {e}")
    values, mapping, unmatched = map_lines_to_columns(result["lines"], columns)
    result["mapped_values"] = values
    result["mapping"] = mapping
    result["unmatched"] = unmatched
    lines_sum = sum(ln["amount_cents"] for ln in result["lines"] if ln.get("amount_cents") is not None)
    result["lines_sum_cents"] = lines_sum
    return result


# --------------------------------------------------------------------------- #
# Exports                                                                      #
# --------------------------------------------------------------------------- #
async def _report_rows(register_id: str, month: Optional[str]):
    reg = await get_register_or_404(register_id)
    columns = [c for c in await get_columns(register_id) if c.get("active", True)]
    entries = await list_entries(register_id, month)
    return reg, columns, entries


@api_router.get("/registers/{register_id}/export/excel")
async def export_excel(register_id: str, month: Optional[str] = None):
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    reg, columns, entries = await _report_rows(register_id, month)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Registro"

    headers = ["Data", "Totale Corrispettivi"] + [c["display_name"] for c in columns]
    ws.append(headers)
    bold = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="0F172A")
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col_idx)
        c.font = bold
        c.fill = fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border

    money_fmt = '#,##0.00 "€"'
    totals = {"total": 0}
    for col in columns:
        totals[col["id"]] = 0

    for e in entries:
        row = [e.get("business_date", "")]
        tr = e.get("total_receipts_cents")
        row.append((tr or 0) / 100)
        totals["total"] += tr or 0
        for col in columns:
            amt = (e.get("values") or {}).get(col["id"])
            row.append((amt or 0) / 100 if amt is not None else None)
            totals[col["id"]] += amt or 0
        ws.append(row)

    # totals row
    trow = ["TOTALI", totals["total"] / 100] + [totals[c["id"]] / 100 for c in columns]
    ws.append(trow)

    last_row = ws.max_row
    for r in range(2, last_row + 1):
        for cc in range(2, len(headers) + 1):
            cell = ws.cell(row=r, column=cc)
            cell.number_format = money_fmt
            cell.border = border
        ws.cell(row=r, column=1).border = border
    for cc in range(1, len(headers) + 1):
        ws.cell(row=last_row, column=cc).font = Font(bold=True)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 20
    from openpyxl.utils import get_column_letter
    for i in range(3, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16

    # Controlli sheet
    ws2 = wb.create_sheet("Controlli")
    ws2.append(["Data", "Totale Azzeramento", "Somma Reparti", "Differenza", "Stato"])
    for cc in range(1, 6):
        c = ws2.cell(row=1, column=cc)
        c.font = bold
        c.fill = fill
    for e in entries:
        ws2.append([
            e.get("business_date", ""),
            (e.get("total_receipts_cents") or 0) / 100,
            (e.get("departments_total_cents") or 0) / 100,
            (e.get("difference_cents") or 0) / 100,
            e.get("reconciliation_status", ""),
        ])
    for r in range(2, ws2.max_row + 1):
        for cc in range(2, 5):
            ws2.cell(row=r, column=cc).number_format = money_fmt
    for col_letter in ["A", "B", "C", "D", "E"]:
        ws2.column_dimensions[col_letter].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = _safe_filename(reg["name"], month, "xlsx")
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


@api_router.get("/registers/{register_id}/export/pdf")
async def export_pdf(register_id: str, month: Optional[str] = None):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    reg, columns, entries = await _report_rows(register_id, month)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=10 * mm, bottomMargin=10 * mm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("t", parent=styles["Title"], fontSize=16, spaceAfter=2)
    sub_style = ParagraphStyle("s", parent=styles["Normal"], fontSize=10, textColor=colors.HexColor("#475569"))
    head_style = ParagraphStyle("h", parent=styles["Normal"], fontSize=7.5, textColor=colors.white,
                                alignment=1, leading=9)
    cell_style = ParagraphStyle("c", parent=styles["Normal"], fontSize=8, alignment=2)
    first_cell = ParagraphStyle("fc", parent=styles["Normal"], fontSize=8, alignment=0)

    period = ""
    if month and "-" in month:
        y, m = month.split("-")
        period = f"{MONTHS_IT[int(m) - 1]} {y}"
    elif reg.get("year"):
        period = str(reg["year"])
    elements = [Paragraph(f"REGISTRO DEI CORRISPETTIVI", title_style)]
    sub = reg["name"]
    if reg.get("business_name"):
        sub += f" — {reg['business_name']}"
    if period:
        sub += f" — {period}"
    elements.append(Paragraph(sub, sub_style))
    elements.append(Spacer(1, 6))

    headers = ["Data", "Totale Corrispettivi"] + [c["display_name"] for c in columns]
    header_row = [Paragraph(h, head_style) for h in headers]
    data = [header_row]
    totals = {"total": 0}
    for col in columns:
        totals[col["id"]] = 0
    for e in entries:
        tr = e.get("total_receipts_cents")
        totals["total"] += tr or 0
        row = [Paragraph(e.get("business_date", ""), first_cell),
               Paragraph(format_cents(tr) if tr is not None else "", cell_style)]
        for col in columns:
            amt = (e.get("values") or {}).get(col["id"])
            totals[col["id"]] += amt or 0
            row.append(Paragraph(format_cents(amt) if amt is not None else "", cell_style))
        data.append(row)
    total_row = [Paragraph("TOTALI", first_cell),
                 Paragraph(format_cents(totals["total"]), cell_style)]
    for col in columns:
        total_row.append(Paragraph(format_cents(totals[col["id"]]), cell_style))
    data.append(total_row)

    page_width = landscape(A4)[0] - 20 * mm
    n = len(headers)
    first_w = 20 * mm
    rest = (page_width - first_w) / (n - 1)
    col_widths = [first_w] + [rest] * (n - 1)

    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F8FAFC")]),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E2E8F0")),
        ("LINEABOVE", (0, -1), (-1, -1), 1, colors.HexColor("#0F172A")),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    elements.append(table)
    doc.build(elements)
    buf.seek(0)
    fname = _safe_filename(reg["name"], month, "pdf")
    return StreamingResponse(
        buf, media_type="application/pdf",
        headers={"Content-Disposition": f'inline; filename="{fname}"'},
    )


def _safe_filename(name, month, ext):
    base = re.sub(r"[^a-zA-Z0-9_-]+", "_", name or "registro")
    if month:
        base += f"_{month}"
    return f"{base}.{ext}"


app.include_router(api_router)
app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("startup")
async def migrate_zreports_on_startup():
    """Safe, idempotent migration: every existing DailyEntry that has a
    total_receipts_cents but no z_reports gets a single azzeramento created from
    that total. Totals stay identical, no data is lost."""
    try:
        query = {"total_receipts_cents": {"$ne": None},
                 "$or": [{"z_reports": {"$exists": False}}, {"z_reports": None}, {"z_reports": []}]}
        migrated = 0
        async for e in db.entries.find(query):
            total = e.get("total_receipts_cents")
            if total is None:
                continue
            z = [{
                "id": str(uuid.uuid4()),
                "amount_cents": int(total),
                "detected_date": e.get("business_date"),
                "confidence": 1.0,
                "order_index": 0,
                "created_at": e.get("created_at", now_iso()),
                "updated_at": now_iso(),
            }]
            await db.entries.update_one({"id": e["id"]}, {"$set": {"z_reports": z}})
            migrated += 1
        if migrated:
            logger.info("z_reports migration: initialized %d existing entries", migrated)
    except Exception:
        logger.exception("z_reports migration failed")


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
